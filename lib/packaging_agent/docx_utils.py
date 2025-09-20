"""Utilities for assembling DOCX documents."""
from __future__ import annotations

import io
from datetime import datetime
from typing import Dict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _add_cover_page(document: Document, project_id: str) -> None:
    """Add a simple cover page with project metadata."""

    title = document.add_heading("Packaging Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph(f"Project ID: {project_id}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    generated = document.add_paragraph(
        f"Generated At: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}"
    )
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_page_break()


def _add_section(document: Document, heading: str, body: str, add_page_break: bool) -> None:
    """Append a section with heading and body text."""

    section_heading = document.add_heading(heading.title(), level=1)
    section_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    document.add_paragraph(body)

    if add_page_break:
        document.add_page_break()


def assemble_docx(project_id: str, sections: Dict[str, str]) -> bytes:
    """Create a DOCX document and return its bytes."""

    document = Document()
    _add_cover_page(document, project_id)

    items = list(sections.items())
    total = len(items)
    for idx, (key, value) in enumerate(items):
        if value is None:
            continue
        heading = key.replace("_", " ")
        add_page_break = idx < total - 1
        _add_section(document, heading, value, add_page_break)

    with io.BytesIO() as buffer:
        document.save(buffer)
        return buffer.getvalue()
